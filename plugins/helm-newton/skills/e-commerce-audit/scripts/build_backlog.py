#!/usr/bin/env python3
"""Build the Newton backlog workbook from scored.json, and verify it after recalculation.

Usage:
    python3 build_backlog.py build scored.json -o backlog.xlsx
    python3 build_backlog.py verify scored.json backlog.xlsx   # run AFTER the xlsx skill's recalc.py

The Backlog sheet holds live formulas (impact band, effort band, rigor weight,
confidence, priority, quadrant, test first, PR/FAQ). Thresholds live on the
Settings sheet. `verify` checks the recalculated cells match score.py exactly.

The "Behavior lane" and "Path" sheets are plain tables: the lane is unscored, so
they hold no formulas. `verify` checks the lane's rows, ranks and order against
scored.json. The Path sheet is a copy of scored.json -> path and is not verified.
"""
import argparse
import json
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, __import__("os").path.dirname(__file__))
import score as S  # noqa: E402

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
BUNDLE_FILL = PatternFill("solid", fgColor="E8EEF6")
CHILD_FONT_BLUE = Font(name=FONT, color="0000FF", italic=True)

COLUMNS = [
    # key, header, width, is_input
    ("id", "ID", 11, True), ("title", "Title", 34, True), ("kind", "Kind", 8, True),
    ("bundle", "Bundle", 10, True), ("cards", "Card(s)", 16, True), ("status", "Status", 9, True),
    ("prev_status", "Prev status", 10, True), ("change", "Change", 10, True),
    ("finding", "Finding", 45, True), ("evidence", "Evidence", 30, True),
    ("origin", "Origin", 16, True), ("timing", "Timing", 14, True),
    ("margin_base_eur", "Margin base €/yr", 13, True), ("effect_low", "Effect low", 9, True),
    ("effect_high", "Effect high", 9, True), ("estimated", "Estimated", 9, True),
    ("impact_eur", "Impact €/yr", 12, False), ("impact_band", "Impact band", 8, False),
    ("newton_days", "Newton days", 8, True), ("client_days", "Client days", 8, True),
    ("run_cost_month", "Run cost €/mo", 9, True), ("effort_band", "Effort band", 8, False),
    ("rigor", "Rigor grade", 10, True), ("rigor_weight", "Rigor weight", 8, False),
    ("transfer", "Transfer", 8, True), ("est_factor", "Est. factor", 8, False),
    ("confidence", "Confidence", 10, False), ("priority", "Priority", 9, False),
    ("quadrant", "Quadrant", 11, False), ("test_first", "Test first", 8, False),
    ("ai_role", "AI role", 28, True), ("data_tools", "Data/tools", 24, True),
    ("proxy_metric", "Proxy metric", 22, True), ("test_design", "Test design", 22, True),
    ("acceptance_criteria", "Acceptance criteria", 34, True),
    ("safety_critical", "Safety-critical", 9, True), ("pr_faq", "PR/FAQ needed", 9, False),
    ("notes", "Notes", 30, True),
]
COL = {k: get_column_letter(i + 1) for i, (k, *_rest) in enumerate(COLUMNS)}

# Settings sheet cell map
SET = {
    "rigor_range": "Settings!$A$2:$A$6", "weight_range": "Settings!$B$2:$B$6",
    "est": "Settings!$B$9", "thr": "Settings!$B$10", "run": "Settings!$B$11",
    "i1": "Settings!$B$13", "i2": "Settings!$B$14", "i3": "Settings!$B$15", "i4": "Settings!$B$16",
    "e1": "Settings!$B$18", "e2": "Settings!$B$19", "e3": "Settings!$B$20", "e4": "Settings!$B$21",
}


def settings_sheet(wb):
    ws = wb.create_sheet("Settings")
    ws.append(["Rigor grade", "Weight"])
    for g, w in S.RIGOR_WEIGHTS.items():
        ws.append([g, w])
    rows = {
        9: ("Est. factor (impact estimated)", S.EST_FACTOR),
        10: ("Test-first threshold (confidence below)", S.TEST_FIRST_THRESHOLD),
        11: ("Run cost raising effort one band (€/month above)", S.RUN_COST_THRESHOLD),
        12: ("Impact band upper bounds (€/yr, band 1 below first)", None),
        17: ("Effort band upper bounds (person-days, inclusive)", None),
    }
    for r, (label, val) in rows.items():
        ws.cell(r, 1, label)
        if val is not None:
            ws.cell(r, 2, val)
    for i, b in enumerate(S.IMPACT_BOUNDS):
        ws.cell(13 + i, 1, f"Band {i + 1} below"); ws.cell(13 + i, 2, b)
    for i, b in enumerate(S.EFFORT_BOUNDS):
        ws.cell(18 + i, 1, f"Band {i + 1} up to"); ws.cell(18 + i, 2, b)
    ws["D1"] = "=1+1"  # recalc sentinel checked by verify
    ws["C1"] = "Recalc check"
    ws.cell(23, 1, "Blue cells are inputs. Change a threshold here and every Backlog formula follows. "
                   "score.py uses the same defaults; re-run verify after changing them.")
    for row in ws.iter_rows():
        for c in row:
            c.font = BLUE if (c.column == 2 and c.value is not None) else BLACK
    ws["A1"].font = ws["B1"].font = Font(name=FONT, bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 12


def formulas(r, kind, child_rows):
    c = COL
    f = {}
    if kind == "bundle":
        if child_rows:
            rng = f"{c['impact_eur']}{child_rows[0]}:{c['impact_eur']}{child_rows[1]}"
            f["impact_eur"] = f"=ROUND(MAX({rng})+0.5*(SUM({rng})-MAX({rng})),{S.EUR_DP})"
        else:
            f["impact_eur"] = '=""'
    else:
        m, lo, hi = (f"{c[k]}{r}" for k in ("margin_base_eur", "effect_low", "effect_high"))
        f["impact_eur"] = f'=IF(OR({m}="",{lo}="",{hi}=""),"",ROUND({m}*({lo}+{hi})/2,{S.EUR_DP}))'
    if kind == "child":
        return f
    q, st = f"{c['impact_eur']}{r}", f"{c['status']}{r}"
    inactive = f'OR({st}="done",{st}="n/a")'
    f["impact_band"] = (f'=IF(OR({inactive},{q}=""),"",IF({q}<{SET["i1"]},1,IF({q}<{SET["i2"]},2,'
                        f'IF({q}<{SET["i3"]},3,IF({q}<{SET["i4"]},4,5)))))')
    d = f"ROUND({c['newton_days']}{r}+{c['client_days']}{r},{S.DAYS_DP})"
    f["effort_band"] = (f'=IF({inactive},"",MIN(5,IF({d}<={SET["e1"]},1,IF({d}<={SET["e2"]},2,'
                        f'IF({d}<={SET["e3"]},3,IF({d}<={SET["e4"]},4,5))))'
                        f'+IF(N({c["run_cost_month"]}{r})>{SET["run"]},1,0)))')
    w = f"{c['rigor']}{r}"
    f["rigor_weight"] = f'=IF({w}="","",INDEX({SET["weight_range"]},MATCH({w},{SET["rigor_range"]},0)))'
    f["est_factor"] = f'=IF({c["estimated"]}{r}="Y",{SET["est"]},1)'
    rw, t, e = f"{c['rigor_weight']}{r}", f"{c['transfer']}{r}", f"{c['est_factor']}{r}"
    f["confidence"] = f'=IF(OR({inactive},{rw}="",{t}=""),"",ROUND({rw}*{t}*{e},{S.CONF_DP}))'
    ib, eb, cf = f"{c['impact_band']}{r}", f"{c['effort_band']}{r}", f"{c['confidence']}{r}"
    f["priority"] = f'=IF(OR({ib}="",{eb}="",{cf}=""),"",{ib}*{cf}/{eb})'
    f["quadrant"] = (f'=IF({st}="done","Done",IF({st}="n/a","N/A",IF(OR({ib}="",{eb}=""),"",'
                     f'IF(AND({ib}>=3,{eb}<=2),"Quick win",IF(AND({ib}>=3,{eb}>=3),"Big bet",'
                     f'IF({eb}<=2,"Fill-in","Skip"))))))')
    f["test_first"] = f'=IF({cf}="","",IF({cf}<{SET["thr"]},"Y","N"))'
    f["pr_faq"] = f'=IF({eb}="","",IF({eb}>=4,"Y","N"))'
    return f


def cell_value(fnd, key):
    v = fnd.get(key)
    if key == "cards":
        return ", ".join(v or [])
    if key in ("estimated", "safety_critical"):
        return "Y" if v else "N"
    if v is None:
        return None
    return v


def backlog_sheet(wb, findings):
    ws = wb.active
    ws.title = "Backlog"
    for i, (_k, head, width, _inp) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, i, head)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    # pre-compute child row ranges (active children are first, contiguous)
    row_of = {f["id"]: 2 + i for i, f in enumerate(findings)}
    child_range = {}
    for f in findings:
        if f["kind"] == "bundle":
            act = [row_of[c["id"]] for c in findings if c["kind"] == "child" and c.get("bundle") == f["id"]
                   and c["status"] not in S.INACTIVE and c.get("impact_eur") is not None]
            child_range[f["id"]] = (min(act), max(act)) if act else None
    for f in findings:
        r = row_of[f["id"]]
        fm = formulas(r, f["kind"], child_range.get(f["id"]))
        for i, (k, _h, _w, is_input) in enumerate(COLUMNS, start=1):
            cell = ws.cell(r, i)
            if k in fm:
                cell.value, cell.font = fm[k], BLACK
            elif is_input:
                # bundle rows: impact inputs come from children; rigor/transfer/estimated are copied values
                if f["kind"] == "bundle" and k in ("margin_base_eur", "effect_low", "effect_high"):
                    continue
                if f["kind"] == "child" and k in ("newton_days", "client_days", "run_cost_month"):
                    continue
                cell.value = cell_value(f, k)
                cell.font = BLUE
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if f["kind"] == "bundle":
                cell.fill = BUNDLE_FILL
        if f["kind"] == "bundle":
            ws[f"{COL['rigor']}{r}"].comment = Comment(
                "Rigor, transfer and estimated are copied from the child with the largest effect (score.py).",
                "e-commerce-audit")
    nf = {"margin_base_eur": "#,##0 €", "impact_eur": "#,##0 €", "run_cost_month": "#,##0 €",
          "effect_low": "0.0%", "effect_high": "0.0%", "confidence": "0.00", "priority": "0.00",
          "rigor_weight": "0.0", "transfer": "0.0", "est_factor": "0.0"}
    for k, fmt in nf.items():
        for r in range(2, len(findings) + 2):
            ws[f"{COL[k]}{r}"].number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(findings) + 1}"


def simple_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        cells = [row.get(h) if isinstance(row, dict) else row[i] for i, h in enumerate(headers)]
        ws.append(["; ".join(map(str, c)) if isinstance(c, (list, tuple)) else
                   json.dumps(c, ensure_ascii=False) if isinstance(c, dict) else c for c in cells])
    for i, h in enumerate(headers, start=1):
        c = ws.cell(1, i)
        c.font, c.fill = BOLD, HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(60, len(h) + 18))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BLACK
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


BEHAVIOR_HEADERS = ["lane_rank", "id", "kind", "title", "status", "prev_status", "change", "grade", "effort",
                    "observation", "hypothesis", "evidence", "source", "source_url", "truth_source", "measure",
                    "test_design", "acceptance_criteria", "origin", "notes"]


def path_rows(scored):
    """Path comparison rows (from path_metrics.py) as flat dicts: metric + one column per walk."""
    rows = scored.get("path") or []
    names = []
    for row in rows:
        for k in row.get("values", {}):
            if k not in names:
                names.append(k)
    names.sort(key=lambda n: (not n.startswith("client"), n))
    return names, [{"metric": r.get("metric"), **r.get("values", {})} for r in rows]


def build(scored, out):
    wb = Workbook()
    backlog_sheet(wb, scored["findings"])
    simple_sheet(wb, "Behavior lane", BEHAVIOR_HEADERS, scored.get("behavior") or [])
    names, rows = path_rows(scored)
    simple_sheet(wb, "Path", ["metric"] + names, rows)
    simple_sheet(wb, "Legal lane", ["id", "rule", "checked", "finding", "evidence", "action"],
                 scored.get("legal", []))
    simple_sheet(wb, "Baseline", ["metric", "value", "date", "method", "previous_value", "previous_date", "change"],
                 scored.get("baseline", []))
    comps = scored.get("competitors", [])
    names = sorted({k for row in comps for k in row.get("values", {})}, key=lambda n: (n != "client", n))
    simple_sheet(wb, "Competitors", ["feature"] + names,
                 [{"feature": r["feature"], **r.get("values", {})} for r in comps])
    simple_sheet(wb, "Not checked", ["item", "reason", "effect"], scored.get("not_checked", []))
    simple_sheet(wb, "Sources", ["card", "source", "url", "rigor"], scored.get("sources", []))
    meta = scored.get("meta", {})
    ws = simple_sheet(wb, "About", ["field", "value"], [{"field": k, "value": str(v)} for k, v in meta.items()])
    ws.append(["generated_by", "e-commerce-audit build_backlog.py"])
    settings_sheet(wb)
    wb.save(out)
    return out


def verify(scored, path, tol=1e-6):
    wb = load_workbook(path, data_only=True)
    ws = wb["Backlog"]
    headers = {ws.cell(1, i).value: i for i in range(1, ws.max_column + 1)}
    keymap = {h: k for k, h, *_ in COLUMNS}
    checks = ["impact_eur", "impact_band", "effort_band", "rigor_weight", "confidence", "priority",
              "quadrant", "test_first", "pr_faq"]
    problems = []
    for r, f in enumerate(scored["findings"], start=2):
        if ws.cell(r, headers["ID"]).value != f["id"]:
            problems.append(f"row {r}: ID mismatch")
            continue
        for k in checks:
            head = next(h for h, kk in keymap.items() if kk == k)
            got = ws.cell(r, headers[head]).value
            exp = f.get(k)
            if f["kind"] == "child" and k != "impact_eur":
                continue
            if k in ("test_first", "pr_faq") and exp is not None:
                exp = "Y" if exp else "N"
            if exp is None:
                if got not in (None, ""):
                    problems.append(f"{f['id']}.{k}: expected blank, got {got!r}")
            elif isinstance(exp, (int, float)):
                if not isinstance(got, (int, float)) or abs(got - exp) > tol * max(1, abs(exp)):
                    problems.append(f"{f['id']}.{k}: expected {exp}, got {got!r}")
            elif got != exp:
                problems.append(f"{f['id']}.{k}: expected {exp!r}, got {got!r}")
    if wb["Settings"]["D1"].value != 2:
        return ["no cached values: run the xlsx skill's recalc.py before verify"]
    problems += verify_behavior(scored, wb)
    return problems


def verify_behavior(scored, wb):
    """The lane has no formulas: check that every row, its rank and its order survived the build."""
    lane = scored.get("behavior") or []
    if "Behavior lane" not in wb.sheetnames:
        return ["Behavior lane sheet is missing"] if lane else []
    ws = wb["Behavior lane"]
    head = {ws.cell(1, i).value: i for i in range(1, ws.max_column + 1)}
    problems = []
    if ws.max_row - 1 != len(lane) and not (len(lane) == 0 and ws.max_row == 1):
        problems.append(f"Behavior lane: expected {len(lane)} rows, got {ws.max_row - 1}")
    for r, b in enumerate(lane, start=2):
        for k in ("id", "kind", "status", "grade", "effort", "lane_rank"):
            got, exp = ws.cell(r, head[k]).value, b.get(k)
            if got != exp:
                problems.append(f"Behavior lane row {r} {k}: expected {exp!r}, got {got!r}")
    return problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build"); b.add_argument("scored"); b.add_argument("-o", "--out", default="backlog.xlsx")
    v = sub.add_parser("verify"); v.add_argument("scored"); v.add_argument("xlsx")
    a = ap.parse_args()
    with open(a.scored, encoding="utf-8") as fh:
        scored = json.load(fh)
    if a.cmd == "build":
        print(json.dumps({"out": build(scored, a.out), "rows": len(scored["findings"])}))
    else:
        problems = verify(scored, a.xlsx)
        print(json.dumps({"status": "ok" if not problems else "mismatch", "problems": problems[:50],
                          "total_problems": len(problems)}, ensure_ascii=False))
        sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
