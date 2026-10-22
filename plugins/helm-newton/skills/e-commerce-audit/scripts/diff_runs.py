#!/usr/bin/env python3
"""Compare a new audit with the previous backlog workbook.

Usage:
    python3 diff_runs.py previous_backlog.xlsx scored.json -o scored.json

Adds prev_status / change to each finding and each behavior-lane item, and
previous_value / change to each baseline row, and lists IDs that disappeared. Run after score.py and before
build_backlog.py. Rewrites scored.json in place by default.
"""
import argparse
import json

from openpyxl import load_workbook

RANK = {"missing": 0, "partial": 1, "done": 2, "n/a": None}


def _heads(ws):
    """Lower-cased header -> column. 'Status (done/partial/missing)' is also indexed as 'status'."""
    out = {}
    for i in range(1, ws.max_column + 1):
        v = ws.cell(1, i).value
        if v is None:
            continue
        h = str(v).strip().lower()
        out.setdefault(h, i)
        out.setdefault(h.split(" (")[0].strip(), i)
    return out


def read_previous(path):
    wb = load_workbook(path, data_only=True)
    prev_status, prev_baseline = {}, {}
    ws = wb["Backlog"]
    head = _heads(ws)
    if "id" not in head or "status" not in head:
        raise SystemExit("previous workbook: Backlog sheet needs ID and Status columns")
    for r in range(2, ws.max_row + 1):
        fid = ws.cell(r, head["id"]).value
        if fid:
            st = ws.cell(r, head["status"]).value
            prev_status[str(fid).strip()] = str(st).strip().lower() if st else None
    if "Behavior lane" in wb.sheetnames:      # absent in workbooks built before plugin 0.2.0
        ws = wb["Behavior lane"]
        head = _heads(ws)
        if {"id", "status"} <= set(head):
            for r in range(2, ws.max_row + 1):
                bid = ws.cell(r, head["id"]).value
                if bid:
                    st = ws.cell(r, head["status"]).value
                    prev_status[str(bid).strip()] = str(st).strip().lower() if st else None
    if "Baseline" in wb.sheetnames:
        ws = wb["Baseline"]
        head = _heads(ws)
        if not {"metric", "value"} <= set(head):
            return prev_status, prev_baseline
        for r in range(2, ws.max_row + 1):
            metric = ws.cell(r, head["metric"]).value
            if metric:
                prev_baseline[metric] = (ws.cell(r, head["value"]).value,
                                         ws.cell(r, head["date"]).value if "date" in head else None)
    return prev_status, prev_baseline


def status_change(old, new):
    if old is None:
        return "new"
    if old == new:
        return "same"
    a, b = RANK.get(old), RANK.get(new)
    if a is None or b is None:
        return f"{old} -> {new}"
    return "improved" if b > a else "regressed"


def apply(scored, prev_status, prev_baseline):
    seen = set()
    rows = list(scored["findings"]) + list(scored.get("behavior", []))
    for f in rows:
        old = prev_status.get(f["id"])
        f["prev_status"] = old
        f["change"] = status_change(old, f["status"])
        seen.add(f["id"])
    for row in scored.get("baseline", []):
        if row["metric"] in prev_baseline:
            pv, pd = prev_baseline[row["metric"]]
            row["previous_value"], row["previous_date"] = pv, pd
            if isinstance(pv, (int, float)) and isinstance(row.get("value"), (int, float)):
                row["change"] = round(row["value"] - pv, 6)
    # A lane that was not run this time (empty list) removes nothing: a lane-only re-run keeps the scored IDs.
    lane_prefixes = ("BEH-", "PATH-", "NEWB-")
    ran_scored, ran_lane = bool(scored["findings"]), bool(scored.get("behavior"))
    removed = sorted(i for i in set(prev_status) - seen
                     if (ran_lane if str(i).startswith(lane_prefixes) else ran_scored))
    scored["diff"] = {
        "removed_ids": removed,
        "improved": [f["id"] for f in rows if f["change"] == "improved"],
        "regressed": [f["id"] for f in rows if f["change"] == "regressed"],
        "new": [f["id"] for f in rows if f["change"] == "new"],
    }
    return scored


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("previous_xlsx")
    ap.add_argument("scored")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    with open(a.scored, encoding="utf-8") as fh:
        scored = json.load(fh)
    ps, pb = read_previous(a.previous_xlsx)
    scored = apply(scored, ps, pb)
    out = a.out or a.scored
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(scored, fh, ensure_ascii=False, indent=2)
    print(json.dumps({"out": out, **{k: len(v) for k, v in scored["diff"].items()}}))


if __name__ == "__main__":
    main()
